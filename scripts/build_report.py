"""Generate the current report from the retained real benchmark manifest."""
import json
from pathlib import Path
from xml.sax.saxutils import escape
from docx import Document
from docx.shared import Inches, Pt
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
DATA = ROOT / "test_sets"

SECTIONS = [
    ("Abstract", "IEAI-ALP is an access-governance research application with a web frontend, Python API, role-mining engine, approval-backed policy changes, audit history, and an AWS serverless deployment option. Five published access-control matrices provide real evaluation inputs. The project preserves the original anonymized assignments and labels six extra CSV fields as computed graph features, never as original employee or ERP attributes."),
    ("Problem and boundaries", "Flat user-to-permission grants are difficult to inspect and maintain. The application groups access profiles, discovers candidate roles, measures representation cost, highlights structurally unusual grants, and lets an administrator review, revoke, or restore a grant with policy-version checks. These are benchmark-policy experiments. They do not provision a live ERP, prove business-role correctness, identify threats, or establish patentability."),
    ("Datasets and provenance", "The test_sets directory contains Americas Small, APJ, Healthcare, Domino, and Firewall2. These are distinct real-world access matrices in the HP Labs role-mining benchmark collection, representing separate source systems/domains; they are not five time partitions. The raw two-column source files are retained alongside import-ready CSVs. Each CSV has the source user and permission IDs plus six deterministic statistics calculated only from that matrix. A builder verifies published dimensions and records SHA-256 hashes in manifest.json. Dataset source links, counts, feature formulas and attribution notes are in test_sets/README.md."),
    ("Feature columns", "The additional columns are derived_user_assignment_count, derived_permission_user_count, derived_user_access_density, derived_permission_prevalence, derived_user_degree_percentile, and derived_permission_degree_percentile. They are computed from the same source edge list, validated at ingestion, and recomputed for application analytics. They help describe grant volume, rarity and prevalence. They are not demographics, timestamps, approval outcomes, actions, or labels."),
    ("Role discovery and optimization", "The engine starts from exact user-permission assignments, groups identical access sets, factors shared permission support signatures using a bounded heuristic, and compares direct, cohort and factored representations using unit WSC: roles plus user-role assignments plus role-permission assignments. It checks that the selected representation reconstructs every original grant without missing or extra grants. This check establishes exact reconstruction for the chosen result, not global optimality."),
    ("User profiling and risk review", "Profiles show observed assignment counts, permission density, rare-grant counts, mined memberships and degree percentiles. Permission catalogs show actual support/prevalence. Review scores use explicitly documented structural features and are cues for a person to inspect; no threat labels or business impact are inferred. Since these benchmark matrices are static, temporal behavior and approval-delay prediction are unavailable."),
    ("Policy changes and reinforcement learning", "Administrators can retain, revoke or restore grants. Expected-version checks prevent stale changes; parent versions, actors and audit events record outcomes. Authorization is checked against the active mined policy. Offline terminal Q-learning is limited to actual administrator feedback, requires at least five genuine reviews across two users, and abstains when evidence is insufficient. No synthetic reviews, rewards, or benchmark-based temporal outcomes are included."),
    ("Cloud architecture", "The SAM/CloudFormation definition configures API Gateway, Lambda, Cognito, private S3, DynamoDB, IAM and CloudWatch Logs; AWS Budgets provides alerts and Access Analyzer validates IAM policy documents. S3 stores uploaded source files and computed artifacts, Lambda runs the API and analysis, Cognito authenticates the administrator, and DynamoDB stores metadata and policy/audit state. The browser can remain local. Deployment and usage depend on the account; free-tier credits are not guaranteed and a budget alert does not cap charges."),
    ("Reproducibility", "Run python scripts/build_test_sets.py to rebuild the five CSVs and manifest from the retained raw matrices. Import a CSV through the local application or AWS upload workflow. The manifest records source hashes and expected dimensions. The source data collection and repository are cited in test_sets/README.md. Review dataset terms before redistribution; the mirror does not state a separate dataset license."),
    ("Portfolio positioning", "This project demonstrates data provenance, feature engineering, sparse access analytics, role discovery, API/backend design, policy governance, auditability, and serverless AWS integration. A defensible CV description should state the exact benchmark names and that the extra CSV fields are derived. Novelty or patentability requires separate prior-art and legal review."),
]


def main():
    manifest = json.loads((DATA / "manifest.json").read_text(encoding="utf-8"))
    for item in manifest["datasets"].values():
        if not (DATA / item["csv_file"]).exists() or not (DATA / item["source_file"]).exists():
            raise FileNotFoundError("Dataset manifest points to a missing file")
    counts = [
        ["Benchmark", "Users", "Permissions", "Assignments"],
        *[[name.replace("_", " ").title(), f"{v['users']:,}", f"{v['permissions']:,}", f"{v['assignments']:,}"]
          for name, v in manifest["datasets"].items()],
    ]
    md = "# IEAI-ALP: ERP Access Intelligence & Adaptive Least Privilege\n\n"
    md += "Computer Architecture and Organization project report\n\n"
    for title, body in SECTIONS:
        md += f"## {title}\n\n{body}\n\n"
    md += "## Dataset dimensions\n\n| Benchmark | Users | Permissions | Assignments |\n|---|---:|---:|---:|\n"
    for row in counts[1:]:
        md += "| " + " | ".join(row) + " |\n"
    md += "\n## References\n\n"
    md += "- HP Labs, [Fast Exact and Heuristic Methods for Role Minimization Problems](https://shiftleft.com/mirrors/www.hpl.hp.com/techreports/2008/HPL-2008-33.html).\n"
    md += "- [Public benchmark matrix repository](https://github.com/carblu/rolemining/tree/main/datasets).\n"
    md += "- [Published benchmark dimensions and source-domain summary](https://www.mdpi.com/2073-8994/16/8/1008).\n"
    md += "- [AWS SAM documentation](https://docs.aws.amazon.com/serverless-application-model/latest/developerguide/what-is-sam.html).\n"
    (DOCS / "PROJECT_REPORT.md").write_text(md, encoding="utf-8")

    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(.72)
    sec.bottom_margin = Inches(.72)
    doc.styles["Normal"].font.name = "Aptos"
    doc.styles["Normal"].font.size = Pt(10)
    doc.add_heading("IEAI-ALP", 0)
    subtitle = doc.add_paragraph("ERP Access Intelligence & Adaptive Least Privilege")
    subtitle.style = "Subtitle"
    doc.add_paragraph("Computer Architecture and Organization project report")
    for title, body in SECTIONS:
        doc.add_heading(title, level=1)
        doc.add_paragraph(body)
    doc.add_heading("Dataset dimensions", level=1)
    table = doc.add_table(rows=0, cols=4)
    table.style = "Light Shading Accent 1"
    for row in counts:
        cells = table.add_row().cells
        for cell, value in zip(cells, row):
            cell.text = value
    doc.add_heading("References", level=1)
    for title, url in [
        ("HP Labs report", "https://shiftleft.com/mirrors/www.hpl.hp.com/techreports/2008/HPL-2008-33.html"),
        ("Benchmark matrix repository", "https://github.com/carblu/rolemining/tree/main/datasets"),
        ("Published benchmark dimensions", "https://www.mdpi.com/2073-8994/16/8/1008"),
        ("AWS SAM", "https://docs.aws.amazon.com/serverless-application-model/latest/developerguide/what-is-sam.html"),
    ]:
        doc.add_paragraph(f"{title}: {url}", style="List Bullet")
    doc.save(DOCS / "PROJECT_REPORT.docx")

    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="ReportTitle", parent=styles["Title"], alignment=TA_CENTER, textColor=colors.HexColor("#132b45"), spaceAfter=8))
    story = [Paragraph("IEAI-ALP", styles["ReportTitle"]), Paragraph("ERP Access Intelligence &amp; Adaptive Least Privilege", styles["Heading2"]), Paragraph("Computer Architecture and Organization project report", styles["Normal"]), Spacer(1, 14)]
    for title, body in SECTIONS:
        story.extend([Paragraph(escape(title), styles["Heading2"]), Paragraph(escape(body), styles["BodyText"]), Spacer(1, 8)])
    story.extend([Paragraph("Dataset dimensions", styles["Heading2"]), Spacer(1, 6)])
    table = Table(counts, repeatRows=1, hAlign="LEFT")
    table.setStyle(TableStyle([
        ("BACKGROUND", (0,0), (-1,0), colors.HexColor("#173653")),
        ("TEXTCOLOR", (0,0), (-1,0), colors.white),
        ("GRID", (0,0), (-1,-1), .35, colors.HexColor("#cfdae5")),
        ("ROWBACKGROUNDS", (0,1), (-1,-1), [colors.white, colors.HexColor("#f1f5f9")]),
        ("PADDING", (0,0), (-1,-1), 6),
    ]))
    story.extend([table, Spacer(1, 14), Paragraph("References", styles["Heading2"])])
    for title, url in [
        ("HP Labs report", "https://shiftleft.com/mirrors/www.hpl.hp.com/techreports/2008/HPL-2008-33.html"),
        ("Benchmark matrix repository", "https://github.com/carblu/rolemining/tree/main/datasets"),
        ("Published benchmark dimensions", "https://www.mdpi.com/2073-8994/16/8/1008"),
        ("AWS SAM", "https://docs.aws.amazon.com/serverless-application-model/latest/developerguide/what-is-sam.html"),
    ]:
        story.extend([Paragraph(f'<link href="{escape(url)}" color="#245f91">{escape(title)}</link>', styles["BodyText"]), Spacer(1, 4)])
    SimpleDocTemplate(str(DOCS / "PROJECT_REPORT.pdf"), pagesize=letter, rightMargin=.7*inch, leftMargin=.7*inch, topMargin=.65*inch, bottomMargin=.65*inch).build(story)


if __name__ == "__main__":
    main()
