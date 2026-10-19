"""Create the user's Phase-I report from the supplied PDF and Word template."""
from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'docs'
ASSETS = OUT / 'template_report_assets'
TARGET = OUT / 'AWS_Cloud_Project_Report.docx'

TITLE = 'Intelligent Permission Optimization Framework for University ERP Systems Using Role Mining Analytics'
TEAM = [('Ansh Goel', '24BIT0258'), ('Agastya Pradhan', '24BIT0187'), ('Varun Verma', '24BIT0208')]

def draw_diagrams():
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    from matplotlib.patches import FancyBboxPatch
    def base():
        fig, ax = plt.subplots(figsize=(9, 9))
        ax.set_xlim(0, 10); ax.set_ylim(0, 10); ax.axis('off')
        return fig, ax
    def box(ax, x, y, w, h, title, body, color='#edf4fb'):
        ax.add_patch(FancyBboxPatch((x,y),w,h,boxstyle='round,pad=0.08',facecolor=color,edgecolor='#35536e',linewidth=1.2))
        ax.text(x+w/2,y+h-.22,title,ha='center',va='top',fontsize=10.5,fontweight='bold',color='#173653')
        ax.text(x+w/2,y+h/2-.15,body,ha='center',va='center',fontsize=9.5,linespacing=1.45,color='#172b3b')
    def arrow(ax,a,b,label='',offset=(0,0)):
        ax.annotate('',xy=b,xytext=a,arrowprops=dict(arrowstyle='-|>',color='#35536e',lw=1.3))
        if label: ax.text((a[0]+b[0])/2+offset[0],(a[1]+b[1])/2+offset[1],label,fontsize=8.5,ha='center',va='center',bbox=dict(facecolor='white',edgecolor='none',pad=1))
    fig,ax=base()
    box(ax,.4,8.65,9.2,.8,'ADMINISTRATOR / AUDITOR','Browser dashboard | HTTPS | manual approval')
    box(ax,.4,6.95,4.3,1,'PRESENTATION','CloudFront + private S3 assets')
    box(ax,5.3,6.95,4.3,1,'AUTHENTICATION','Cognito user pool / JWT')
    arrow(ax,(2.55,8.65),(2.55,7.95),'assets',(.5,0))
    arrow(ax,(7.45,8.65),(7.45,7.95),'sign in',(.6,0))
    box(ax,.4,5.2,9.2,1,'API GATEWAY + LAMBDA API HANDLER','JWT authorizer | input validation | status and result APIs')
    arrow(ax,(2.55,6.95),(2.55,6.2),'HTTPS',(.6,0))
    arrow(ax,(7.45,6.95),(7.45,6.2),'JWT',(.4,0))
    box(ax,.4,3.2,5.5,1.1,'S3 INPUTS + LAMBDA ANALYTICS','ERP audit CSV / JSON + attributes\nCloudTrail for supported AWS activity')
    box(ax,6.5,3.2,3.1,1.1,'DYNAMODB','Role / job metadata\nPolicy versions / audit')
    arrow(ax,(3.15,5.2),(3.15,4.3),'upload / job',(.85,0))
    arrow(ax,(8.05,5.2),(8.05,4.3),'read / write',(.8,0))
    arrow(ax,(5.9,3.75),(6.5,3.75))
    box(ax,.4,1.4,5.5,1,'IAM POLICY REVIEW','Access Analyzer checks + regression tests\nAdministrator approves candidate changes')
    box(ax,6.5,1.4,3.1,1,'GRAPH EXTENSION','Neptune in VPC\nOptional; separately costed')
    arrow(ax,(3.15,3.2),(3.15,2.4),'candidate policy',(1.05,0))
    arrow(ax,(8.05,3.2),(8.05,2.4))
    box(ax,.4,.05,9.2,.7,'CROSS-CUTTING CONTROLS','IAM execution roles | CloudWatch | SNS | AWS Budgets alerts', '#f3f0df')
    ax.text(5,9.85,'PROPOSED AWS CLOUD ARCHITECTURE',ha='center',fontsize=13,fontweight='bold')
    fig.savefig(ASSETS/'aws_architecture.png',dpi=230,bbox_inches='tight',facecolor='white'); plt.close(fig)
    fig,ax=base()
    blocks=[
        ('1. INPUT AND PROVENANCE','UPA benchmark matrix + authorised ERP attributes / logs\nRecord source, license, schema, hashes and privacy controls'),
        ('2. VALIDATION AND PREPARATION','Deduplicate grants; encode genuine attributes; build sparse matrix\nKeep synthetic features labelled; reserve held-out users / events'),
        ('3. CANDIDATE ROLE DISCOVERY','Attribute clustering, when attributes exist; greedy set-cover\nOptional graph analysis; propose unusual grants for review'),
        ('4. EVALUATION AND POLICY MAPPING','Compute WSC; reconstruct grants; compare consistent baselines\nMap only supported ERP actions to explicit AWS actions / ARNs'),
        ('5. VALIDATION AND ADMINISTRATOR REVIEW','Access Analyzer grammar / security checks + workflow tests\nReject unsafe candidates; record approval and version history'),
        ('6. OUTPUT AND FEEDBACK','Dashboard, role matrix, candidate policy, metrics and audit history\nApproved deployment in test account; rollback on regression')]
    for i,(title,body) in enumerate(blocks):
        y=8.3-i*1.5
        box(ax,.5,y,9,1.05,title,body,'#edf4fb' if i%2==0 else '#f4f2e3')
        if i: arrow(ax,(5,y+1.5),(5,y+1.05))
    ax.text(5,9.8,'PROPOSED COMPLETE SYSTEM ARCHITECTURE',ha='center',fontsize=13,fontweight='bold')
    ax.text(5,.08,'Monitoring and reviewed feedback inform the next analysis cycle.',ha='center',fontsize=9.5)
    fig.savefig(ASSETS/'system_architecture.png',dpi=230,bbox_inches='tight',facecolor='white'); plt.close(fig)

def field(paragraph, instruction, fallback=''):
    run=paragraph.add_run(); begin=OxmlElement('w:fldChar'); begin.set(qn('w:fldCharType'),'begin'); run._r.append(begin)
    r=paragraph.add_run(); ins=OxmlElement('w:instrText'); ins.set(qn('xml:space'),'preserve'); ins.text=instruction; r._r.append(ins)
    r=paragraph.add_run(); sep=OxmlElement('w:fldChar'); sep.set(qn('w:fldCharType'),'separate'); r._r.append(sep)
    paragraph.add_run(fallback)
    r=paragraph.add_run(); end=OxmlElement('w:fldChar'); end.set(qn('w:fldCharType'),'end'); r._r.append(end)

doc=Document(ROOT/'AWS_Cloud_Project_Report_Template.docx')
for child in list(doc._element.body):
    if child.tag != qn('w:sectPr'): doc._element.body.remove(child)
sec=doc.sections[0]
sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Inches(1)
for name in ['Normal','Body Text','List Bullet','List Number','Caption']:
    style=doc.styles[name]; style.font.name='Times New Roman'; style.font.size=Pt(12)
    style.paragraph_format.line_spacing=1.5
    style.paragraph_format.space_after=Pt(6)
for name,size in [('Title',20),('Subtitle',14),('Heading 1',16),('Heading 2',14),('Heading 3',14)]:
    style=doc.styles[name]; style.font.name='Times New Roman'; style.font.size=Pt(size)
    style.font.bold=True; style.font.color.rgb=RGBColor.from_string('173653')
    style.paragraph_format.keep_with_next=True
    style.paragraph_format.line_spacing=1.5
doc.styles['Heading 1'].paragraph_format.page_break_before=True
for p in sec.footer.paragraphs: p.clear()
footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
field(footer,' PAGE ','1')
for p in sec.header.paragraphs: p.clear()
p=sec.header.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=p.add_run('BCSE355L | Cloud Architecture Design | Phase I'); r.font.name='Times New Roman'; r.font.size=Pt(9)
sec.different_first_page_header_footer=True
settings=doc.settings.element
update=OxmlElement('w:updateFields'); update.set(qn('w:val'),'true'); settings.append(update)

list_count=0
def p(text='', style=None):
    global list_count
    if style=='List Number':
        list_count+=1
        z=doc.add_paragraph(f'{list_count}. {text}')
        z.paragraph_format.left_indent=Inches(.25)
        z.paragraph_format.first_line_indent=Inches(-.25)
        return z
    list_count=0
    z=doc.add_paragraph(text,style)
    if style is None: z.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    return z
def h(text,level=1): return doc.add_heading(text,level)
def bullets(items):
    for item in items: p(item,'List Bullet')
def table(headers,rows,widths=None):
    t=doc.add_table(rows=1,cols=len(headers)); t.style='Table Grid'; t.autofit=False
    if widths:
        for col,w in zip(t.columns,widths): col.width=Inches(w)
    for cell,text in zip(t.rows[0].cells,headers): cell.text=text
    for row in rows:
        for cell,text in zip(t.add_row().cells,row): cell.text=str(text)
    for ri,row in enumerate(t.rows):
        for ci,c in enumerate(row.cells):
            if widths: c.width=Inches(widths[ci])
            for para in c.paragraphs:
                para.paragraph_format.line_spacing=1.5
                para.paragraph_format.space_after=Pt(3)
                for r in para.runs: r.font.name='Times New Roman'; r.font.size=Pt(12); r.bold=(ri==0)
            if ri==0:
                shade=OxmlElement('w:shd'); shade.set(qn('w:fill'),'E7EEF5'); c._tc.get_or_add_tcPr().append(shade)
        trpr=row._tr.get_or_add_trPr()
        if ri==0: trpr.append(OxmlElement('w:tblHeader'))
        trpr.append(OxmlElement('w:cantSplit'))
    return t
def caption(text):
    q=p(text,'Caption'); q.alignment=WD_ALIGN_PARAGRAPH.CENTER
    return q

def cover():
    p('PROJECT REPORT','Title').alignment=WD_ALIGN_PARAGRAPH.CENTER
    p('CLOUD ARCHITECTURE DESIGN – AWS CLOUD PROJECT','Subtitle').alignment=WD_ALIGN_PARAGRAPH.CENTER
    p('PHASE I: DESIGN AND IMPLEMENTATION PLAN').alignment=WD_ALIGN_PARAGRAPH.CENTER
    p()
    q=p(TITLE); q.alignment=WD_ALIGN_PARAGRAPH.CENTER
    for r in q.runs: r.bold=True; r.font.size=Pt(17)
    p()
    p('Submitted by:').alignment=WD_ALIGN_PARAGRAPH.CENTER
    for name,reg in TEAM: p(f'{name} — {reg}').alignment=WD_ALIGN_PARAGRAPH.CENTER
    for s in ['Programme / Branch / Year: Not specified in the source report',
              'Course Code and Name: BCSE355L – Cloud Architecture Design',
              'Institution: Vellore Institute of Technology, Vellore',
              'Course Instructor / Guide: Dr. Priya V',
              'Team Number: Not specified in the source report',
              'GitHub Repository: https://github.com/Ansh-bit948/AWS-project',
              'Date of Submission: 31/07/2026 (as recorded in the source PDF)']:
        q=p(s); q.paragraph_format.line_spacing=1.15

ABSTRACT='''University Enterprise Resource Planning (ERP) platforms manage access to academic records and administrative services. As responsibilities change, individually maintained permissions can accumulate redundant privileges and overlapping roles. This project proposes an Intelligent Permission Optimization Framework combining role-mining analytics with an AWS governance workflow for compact, interpretable access policies.

The method accepts user–permission matrices, authorised organisational attributes, and audit records. It prepares sparse representations, applies attribute-guided clustering where genuine attributes exist, and uses greedy role discovery. Weighted Structural Complexity (WSC), permission reconstruction, role counts, and administrator review provide the evaluation framework. AWS policy candidates require explicit mappings from application permissions to supported actions and resources.

The architecture proposes Amazon S3 for artifacts, AWS Lambda for analysis, DynamoDB for policy history, API Gateway for access, and Cognito for authentication. CloudWatch, SNS, and AWS Budgets support monitoring and alerts. Neptune is an optional, separately costed graph extension. IAM Access Analyzer supports scoped policy checks alongside workflow tests and approval.

Phase-I outputs comprise the problem formulation, literature synthesis, objectives, architecture, dataset plan, and roadmap. Clustering is reported in progress and policy generation planned. Runtime, cost, WSC reductions, and security outcomes remain targets because supporting measurements are absent. The expected contribution is a reproducible pipeline connecting access analytics to controlled policy decisions.'''

PAPERS=[
('P1','Molloy et al. (2009)','Evaluating role mining algorithms','WSC-based comparison; candidate-list and complete-state role miners.','Real access matrices and synthetic generators, as described in the source.','Provides a common structural comparison framework.','Structural metrics do not by themselves establish university job-role meaning.','Use WSC alongside reconstruction checks and administrator interpretation.'),
('P2','Lu et al. (2014)','An optimization framework for role mining','Integer linear programming for role-mining variants; heuristic alternatives.','Real and synthetic user–permission matrices.','Expresses competing role and assignment objectives explicitly.','Exact optimisation can become expensive; semantic validity still needs separate evidence.','Use bounded heuristics and document runtime and quality trade-offs.'),
('P3','Anderer et al. (2020)','The addRole-EA: A new evolutionary algorithm for the role mining problem','Evolutionary role construction using sparse binary representations.','HP Labs benchmark instances, as reported in the source.','Maintains permission consistency while exploring role structures.','Evolutionary search and operator selection can require substantial computation.','Compare guided candidates with an unguided baseline under the same budget.'),
('P4','Hamid et al. (2017)','Recent trends in role mining algorithms for role-based access control: A systematic review','Taxonomy of role-mining problems and solution families.','Literature review; no new project benchmark.','Connects algorithm choices with operational concerns.','A survey does not directly establish deployment performance for this project.','Include administrator usability and interpretability in evaluation.'),
('P5','Mollakuqe and Dimitrova (2024)','Comparative analysis of identity management, access control, and authorization practices in public and private universities','Empirical comparison of university identity and access management practices.','University survey data, as described in the supplied PDF.','Provides application-domain motivation for account lifecycle governance.','Operational observations do not supply a tested permission-mining engine.','Evaluate mining with real governance workflows and authorised data.'),
('P11','Attribution incomplete in source (2021)','Mining ABAC rules from sparse logs','Source describes statistical reliability measures for incomplete access traces.','Sparse enterprise/education logs claimed; dataset identifiers absent.','Motivates caution when observations cover only part of legitimate usage.','Exact publication, authors, and reported dataset need verification.','Treat absent events as missing evidence; require review before revocation.'),
('P12','Attribution incomplete in source (2026)','ALPS: Automated least-privilege enforcement for securing serverless functions','Source describes static/trace analysis for cloud policy synthesis.','Source claims 8,000+ repositories; no retained evaluation evidence.','Motivates explicit action/resource policy construction.','Citation and repository-scale claim remain unverified.','Validate the publication before relying on its quantitative claims.'),
('P13','Díaz-Rodríguez and Díaz-Pérez, as supplied (2025)','ABAC policy mining using complex network analysis techniques','Source describes bipartite graphs and community-based analysis.','Access logs with injected noise, according to the source.','Motivates graph structure as an additional review signal.','Bibliographic details and noise-performance claims require verification.','Keep anomaly flags separate from automatic deletion of permissions.'),
('P14','Karimi et al., as supplied (2022)','An unsupervised learning / hierarchical clustering based approach for mining ABAC policies','Source describes clustering over subject and resource attributes.','Attributed and synthetic access tuples, according to the source.','Motivates attribute-assisted candidate policy generation.','The source combines alternative titles and venues; exact record is unresolved.','Resolve the citation and assess clustering with labelled attribute provenance.'),
('P15','Attribution incomplete in source (2025)','Comprehensive identity and access management in AWS: Authentication, authorization, and policy control mechanisms','Source describes an AWS governance architecture.','Case studies claimed; no reproducible study identifier supplied.','Motivates integration of authentication, policy controls, and auditing.','Generic group author and ambiguous venue do not identify a verified paper.','Use official AWS documentation for service behaviour pending verification.')]

def literature():
    h('2. Literature Survey / Related Work')
    p('The supplied PDF allocates paper identifiers P1–P5 to Ansh Goel, P6–P10 to Agastya Pradhan, and P11–P15 to Varun Verma. P6–P10 repeat the titles and analyses of P1–P5. Consequently, the source contains 15 numbered slots but only 10 distinct claimed titles; several of the latter have incomplete bibliographic records. The original mapping is retained below without counting duplicates as new research. Older foundational role-mining studies are retained because they underpin the proposed WSC and reconstruction methodology.')
    caption('Table 1. Source paper identifiers and student allocation.')
    rows=[]
    for i,a in enumerate(PAPERS[:5]): rows.append((a[0],a[1]+' — '+a[2],'Ansh Goel'))
    for i,a in enumerate(PAPERS[:5]): rows.append((f'P{i+6}',f'Duplicate of {a[0]}: '+a[2],'Agastya Pradhan'))
    for a in PAPERS[5:]: rows.append((a[0],a[1]+' — '+a[2],'Varun Verma'))
    table(['Paper ID','Publication / source identity','Source allocation'],rows,[.65,4.35,1.5])
    for a in PAPERS:
        ident,authors,title,method,data,adv,lim,imp=a
        h(f'{ident}. {title}',2)
        p(f'{authors}. Bibliographic details are reproduced from the supplied PDF; unresolved entries are identified in the references.')
        p('Problem and method: '+method+' Dataset / evaluation: '+data)
        p('Advantage: '+adv+' Limitation: '+lim)
        p('Project relevance and possible improvement: '+imp)
    h('2.1 Synthesis and Comparison',2)
    p('The foundational studies address structural optimisation, while the university study motivates operational governance. The additional source entries extend the design toward sparse traces, graph analysis, attributes, and AWS integration. These categories support a combined research direction, but they do not establish that earlier methods are operationally useless or that the proposed method is globally optimal. Comparisons must use the same input grants, WSC weights, runtime limits, and reconstruction criteria. Cloud latency, storage, and cost require separate deployment measurements.')
    caption('Table 2. Method comparison and project research gaps.')
    table(['Method / sources','Advantages','Limitations and proposed improvement'],[
        ('Structural optimisation: P1/P6, P2/P7','Quantifies role and edge complexity.','Add reconstruction and human interpretation; do not equate smaller WSC with safer access.'),
        ('Evolutionary mining: P3/P8','Searches alternative exact role structures.','Compare runtime and solution quality with a bounded heuristic.'),
        ('Survey and university practice: P4/P9, P5/P10','Connects role engineering with administrator needs.','Measure review effort using an explicit study protocol.'),
        ('Sparse / graph / attribute methods: P11, P13, P14','Proposed support for incomplete logs and candidate grouping.','Resolve citations; validate noise handling without deleting legitimate rare grants.'),
        ('Cloud policy integration: P12, P15','Proposed link from analytics to enforceable controls.','Resolve citations; implement explicit mappings and scoped policy checks.')],[2,1.8,2.7])
    h('2.2 Individual Research Gap Analysis',2)
    p('Ansh Goel: The P1–P5 source analysis motivates combining structural role optimisation with operational interpretability. The proposed improvement is to supplement WSC with genuine organisational attributes and administrator review. Attribute availability and semantic labels must be demonstrated before claiming business-role alignment.')
    p('Agastya Pradhan: The supplied P6–P10 analysis covers the same five papers and repeats the same gaps. It supports the clustering and governance direction but does not establish an independent set of five additional publications. Five distinct replacement studies and an independently authored comparison are still required to meet the template’s individual literature expectation.')
    p('Varun Verma: The P11–P15 source analysis proposes connecting sparse/graph analytics to AWS policy synthesis and validation. The useful gap is the missing, tested mapping between application grants and deployable cloud permissions. Exact publication records and the scope of each policy check must be resolved before asserting formal security guarantees.')
    h('2.3 Consolidated Research Gap',2)
    p('The project will investigate whether a single reproducible workflow can preserve approved access, reduce structural policy complexity, improve administrator interpretation, and produce explicitly scoped AWS policy candidates. This is an integration and evaluation question. A systematic prior-art comparison and measured experiments are needed before making a stronger novelty or superiority claim.')

def main():
    ASSETS.mkdir(parents=True,exist_ok=True); draw_diagrams(); cover()
    h('Abstract')
    for para in ABSTRACT.split('\n\n'): p(para)
    h('Keywords',2); p('Role mining; least privilege; university ERP; AWS cloud architecture; access governance; Weighted Structural Complexity.')
    h('Table of Contents')
    field(p(),' TOC \\o "1-1" \\h \\z ','Table of contents will be generated by Microsoft Word.')
    h('1. Introduction')
    for title,text in [
        ('1.1 Background','University ERP systems coordinate academic and administrative functions involving students, faculty, finance staff, examination officers, and auditors. Individual user–permission assignments become difficult to maintain as responsibilities change. Role-Based Access Control groups permissions into roles, while role mining seeks candidate structures from existing assignments. Cloud services can support centralised processing, access controls, auditability, and on-demand analysis.'),
        ('1.2 Problem Statement','The source project addresses permission bloat and overlapping roles in a university ERP setting. Administrators need a way to inspect grants, discover compact candidates, and review their consequences. A structural reduction alone cannot determine whether a permission is legitimate. The problem therefore includes data provenance, interpretation, permission preservation, explicit cloud-policy mapping, and accountable approval.'),
        ('1.3 Motivation and Significance','An auditable optimisation workflow could reduce policy-maintenance effort while making risky or redundant assignments easier to review. Molloy et al. (2009) motivate structural comparison, Lu et al. (2014) frame optimisation alternatives, and the university identity-management study cited in the source motivates the application domain. Expected benefits remain hypotheses until evaluated.'),
        ('1.4 Scope','Phase I covers literature synthesis, requirements, measurable objectives, two architecture diagrams, dataset planning, AWS service planning, and an implementation roadmap. Phase II proposes role mining, an authenticated dashboard, policy synthesis, validation, and controlled deployment. Live university integration, automated removal of unobserved permissions, universal least-privilege guarantees, and established research novelty are outside the demonstrated Phase-I scope.'),
        ('1.5 Report Organisation','Sections 2–4 cover related work, objectives, and proposed novelty. Sections 5–8 describe architecture, datasets, AWS planning, and methodology. Sections 9–12 distinguish development status, planned tests, limitations, and team workflow. Section 13 concludes with future work; references and appendices record sources, milestones, and review status.')]:
        h(title,2); p(text)
    literature()
    h('3. Project Objectives')
    p('The following objectives adapt the source targets into testable criteria. They are planned success thresholds, not achieved results.')
    for text in [
        'Develop attribute-guided candidate role discovery where genuine user attributes exist. Evaluate silhouette score (source target ≥ 0.65) and candidate-search reduction (source target ≥ 50%) using a documented baseline and identical evaluation conditions.',
        'Compare policy representations using declared WSC weights and exact reconstruction. Investigate the source target of ≥ 40% WSC reduction without adding or losing approved grants; report any target that is not reached.',
        'Produce explicit AWS action/resource policy candidates for supported mappings. Reject Action: "*" in the tested outputs and measure missed and extra grants against a reviewed reference set; do not revoke permissions solely because they are unobserved.',
        'Deploy and measure a small AWS test workload. Investigate the source target of < 2.5 seconds for a specified input size, recording cold/warm latency, memory, concurrency, repetitions, and actual usage. Treat the < US$1 cost goal as an unverified demo budget.',
        'Integrate IAM Access Analyzer policy validation, defined guardrail checks, and workflow regression tests. Record each check and finding; require administrator approval before applying a changed policy.',
        'Deliver an authenticated administrator dashboard showing dataset provenance, role structures, WSC comparisons, review decisions, policy versions, and operational alerts. Verify successful authorised access and rejection of unauthorised requests.'
    ]: p(text,'List Number')
    h('3.1 Expected Outcomes',2)
    p('Expected deliverables are a validated input schema and provenance manifest, a reproducible mining module, baseline comparison outputs, two labelled architecture diagrams, an authenticated dashboard, policy-validation artifacts, deployment configuration, and an evidence-backed final report. Each result must link to the relevant input, code revision, test protocol, and measured output.')
    h('4. Novelty Summary')
    p('The proposed contribution is an integrated workflow joining role mining, documented attribute provenance, structural evaluation, administrator review, and AWS policy checks. The design intends to retain a trace from each candidate policy back to its source grants and approval record. These are project contributions to be evaluated; the source does not establish that their individual components are new. Stronger claims about optimality, one-to-one campus roles, or comprehensive security proof require evidence beyond Phase I.')
    caption('Table 3. Proposed contribution and demonstration method.')
    table(['Dimension','Proposed difference','Evidence required'],[
        ('Feature / usability','Joint display of role structure, grant provenance, and review decisions.','Administrator task protocol and traceable examples.'),
        ('Algorithm / processing','Attribute-assisted candidates with bounded structural optimisation.','Ablation against the same pipeline without attributes.'),
        ('Cloud integration','Explicit mapping from eligible ERP permissions to AWS policy actions/resources.','Mapping tests, validation responses, and regression results.'),
        ('Security / automation','Approval and version history before policy application.','Denied unsafe changes, successful reviewed changes, and rollback evidence.')],[1.3,2.65,2.55])
    h('5. Proposed System and Architecture')
    h('5.1 System Overview',2)
    p('Administrators upload an authorised access matrix and, where available, real organisational attributes and audit logs. The backend validates the input, generates candidate roles, computes WSC and reconstruction metrics, and exposes results to the dashboard. A separate mapping layer prepares AWS policy candidates only for supported actions and resources. Policy checks and administrator approval precede any deployment. This remains the proposed Phase-II system described by the PDF.')
    h('5.2 Diagram 1 – AWS Cloud Architecture (Mandatory)',2)
    p('Figure 1 redraws the source architecture at readable resolution and clarifies the distinction between application audit data, AWS activity logs, optional graph processing, and policy approval.')
    doc.add_picture(str(ASSETS/'aws_architecture.png'),width=Inches(6.45))
    caption('Figure 1. Proposed AWS Cloud Architecture, adapted from source PDF pp. 21–23.')
    p('Cognito supplies authentication tokens to a browser client. API Gateway authorises calls to the Lambda API, which coordinates object uploads, analysis jobs, metadata, and results. S3 stores inputs and artifacts; DynamoDB stores metadata and policy history. A Lambda analysis component produces role candidates and a policy-review component invokes scoped checks. Neptune is an optional extension requiring a separate network and cost decision. CloudWatch, SNS, and AWS Budgets support operations and alerts.')
    h('5.3 Diagram 2 – Complete System Architecture (Mandatory)',2)
    doc.add_picture(str(ASSETS/'system_architecture.png'),width=Inches(6.45))
    caption('Figure 2. Complete System Architecture, adapted from source PDF pp. 22–24.')
    p('The workflow preserves provenance at ingestion and keeps attributes distinct from permission observations. Mining and evaluation precede cloud-policy mapping. An anomaly signal initiates review rather than proving misuse. Validation failures return candidates for correction, while approved changes retain version and audit records. Operational results inform the next analysis cycle.')
    h('5.4 Workflow / Use Cases',2)
    for s in ['Administrator signs in and selects an authorised dataset.','The application validates schema, identifiers, duplicate grants, and attribute provenance.','The analysis constructs a sparse representation and discovers candidate roles.','The evaluator compares WSC, role counts, and reconstruction with a declared baseline.','An administrator reviews candidate roles and any supported AWS policy mapping.','Policy checks and regression tests run before approval and controlled application.','The dashboard presents results and audit history; monitoring alerts support investigation.']: p(s,'List Number')
    h('5.5 Security, Availability and Scalability Design',2)
    p('The proposed controls include least-privilege execution roles, private S3 storage, authenticated API access, encryption settings, upload limits, and sanitised audit records. Separate analysis permissions from policy-application permissions. Preserve approved versions and define rollback. Lambda concurrency and API throttling bound load; long analyses may need asynchronous jobs. Availability, retry behaviour, backup restoration, and capacity limits must be tested. If Neptune is selected, place it in private VPC subnets with an explicitly designed Lambda connection and assess endpoint or egress requirements.')
    p('AWS CloudTrail records supported AWS account activity; it does not automatically provide all university ERP user actions. The ERP therefore needs its own authorised audit export or application instrumentation. IAM Access Analyzer policy validation checks policy grammar and AWS best practices, while external-access analysis and custom checks have distinct scopes. Passing one check is not proof that all business permissions are necessary or safe (Amazon Web Services, n.d.-a, n.d.-b).')
    h('6. Dataset Description')
    p('The PDF describes a combined benchmark and synthetic ERP dataset, but supplies no downloadable dataset, manifest, generation script, or measured file inventory. The following attributes are recorded as source specifications rather than verified dataset measurements. Benchmark matrices and synthetic ERP records must remain separately identifiable.')
    caption('Table 4. Dataset specification and verification status.')
    table(['Dataset attribute','Details'],[
        ('Dataset name','University ERP Access Log & Attribute Matrix, benchmarked on HP Labs role-mining matrices (source proposal).'),
        ('Source / organisation','Source PDF proposes HP Labs benchmarks and simulated university-domain records; references to KL University do not establish authorised access to real institutional records.'),
        ('URL','No direct dataset download or provenance manifest supplied. The project repository is https://github.com/Ansh-bit948/AWS-project.'),
        ('Dataset size','File size in MB/GB not documented.'),
        ('Records / samples','Source specification: 112,120 access tuples, 4,930 users, and 560 permissions. Counts remain unverified.'),
        ('Features','Seven named raw fields: user ID, department, designation, timestamp, API endpoint, action, and resource ID. Source also proposes a 128-dimensional permission representation and 12 domain features; the transformation is unspecified.'),
        ('Data types / format','Proposed CSV assignment matrices, JSON event logs, and categorical attribute tables.'),
        ('License / restrictions','The source names an academic/synthetic license but supplies no license text. Verify benchmark redistribution terms and institutional authorisation separately.'),
        ('Purpose','Candidate-role discovery, attribute clustering, WSC comparison, policy-mapping tests, and dashboard demonstration.'),
        ('Preprocessing','Schema checks, deduplication, sparse matrix construction, attribute encoding, and documented optional graph features.'),
        ('Privacy / ethics','Use anonymised identifiers; keep synthetic attributes labelled; exclude secrets and personal records without permission. A claim of no personal information requires field-level review.')],[1.65,4.85])
    h('6.1 Data Preparation Plan',2)
    p('Retain raw files and record source links, license terms, schema, dimensions, and hashes. Remove duplicate edges only under a documented rule and report missing values. Preserve original benchmark identities; do not relabel anonymous users as actual departments. Real domain attributes require independent source records. Explain any reduction from 560 permissions to a 128-dimensional representation and reconstructability implications before evaluating it.')
    p('For structural role mining, evaluate exact reconstruction on the input matrix and report this as representation fidelity. For generalisation experiments, separate held-out users or later events where supported, fit encoders and clustering on the training portion only, and predeclare splits and seeds. Static benchmark matrices cannot demonstrate temporal adaptation or workflow safety. Unobserved events are not evidence that a permission is unnecessary.')
    h('7. AWS Services Planning')
    p('All service entries below describe the source design. The supplied PDF contains no deployment identifiers, dated console evidence, or reproducible service tests, so the services are marked Planned rather than Deployed. Team ownership follows the source service table.')
    caption('Table 5. AWS services, integration, status, and ownership.')
    services=[
        ('Amazon S3 / CloudFront','Inputs, artifacts, and dashboard delivery.','Private objects; scoped access; storage/request/egress charges.','Ansh Goel'),
        ('AWS Lambda','API handler and role-analysis execution.','Bound runtime, memory, uploads, retries, and concurrency.','Varun Verma'),
        ('Amazon DynamoDB','Metadata, role state, policy versions, audit history.','Choose capacity mode; assess backup and item-size limits.','Agastya Pradhan'),
        ('Amazon API Gateway','Authenticated dashboard API.','Cognito authorizer; throttling and access logs.','Varun Verma'),
        ('Amazon Cognito','Administrator authentication and tokens.','Configure user flows, token expiry, and administrator access.','Agastya Pradhan'),
        ('Amazon Neptune (optional)','Graph-based role candidate analysis.','Private VPC; ongoing capacity and storage costs require separate approval.','Ansh Goel'),
        ('AWS IAM / Access Analyzer','Execution roles and scoped policy checks.','Explicit actions/resources; distinguish validation from exposure analysis.','Varun Verma'),
        ('AWS CloudTrail','Supported AWS activity and audit records.','Select event types, retention, and data-event scope.','Owner not specified'),
        ('Amazon CloudWatch','Logs, metrics, alarms.','Set retention and avoid sensitive log content.','Ansh Goel'),
        ('Amazon SNS','Notifications for review and operational events.','Scoped publishing; channel-specific charges.','Agastya Pradhan'),
        ('AWS Budgets','Actual/forecast cost threshold alerts.','US$1 is a proposed alert threshold; it does not cap charges.','All members')]
    table(['Service / owner','Purpose and integration','Status / considerations'],[(s+'\n'+o,pu,'Planned. '+co) for s,pu,co,o in services],[1.65,2.1,2.75])
    h('7.1 AWS Deployment and Network Design',2)
    p('The source does not specify a region, account, VPC identifiers, or infrastructure deployment record. Select one documented region after checking service availability and cost. Use separate development configuration and avoid placing secrets in code. The core API/storage design can use managed service endpoints; the optional Neptune branch requires private VPC connectivity and security-group rules. Define the deployment method, resource naming, environment settings, teardown process, and a test account before Phase-II deployment.')
    h('7.2 Cost Estimation',2)
    p('The source proposes a total cost below US$1 but provides neither a calculator estimate nor measured usage. That amount is retained as a goal, not a validated monthly estimate. Record analysis job count, API requests, Lambda GB-seconds, S3 storage/requests, database usage, logging, notifications, and transfer before calculating a region-specific estimate. Verify account-specific Free Tier or credit eligibility rather than copying the fixed allowances in the original PDF.')
    p('Official AWS documentation reviewed on 15 October 2026 describes Neptune Serverless capacity billing and a minimum capacity of 1 NCU. The proposed Neptune-inclusive design therefore does not justify a blanket zero-idle-cost claim. AWS Budgets sends threshold alerts; it is not a general hard spending cap (Amazon Web Services, n.d.-c, n.d.-d). A numeric monthly estimate remains pending a workload, selected region, and pricing calculation.')
    caption('Table 6. Cost worksheet to be completed before deployment.')
    table(['Cost component','Required quantity / assumption','Current estimate'],[
        ('Compute and API','Requests per month, execution duration, memory, transfer.','Not yet calculated.'),
        ('Storage and database','GB-months, read/write volume, backups, retention.','Not yet calculated.'),
        ('Graph extension','Neptune capacity-hours, storage, I/O, network requirements.','Separate optional estimate required.'),
        ('Operations','Log ingestion, metrics, alarms, SNS delivery, audit events.','Not yet calculated.'),
        ('Monthly total','Selected region and dated official pricing inputs.','Pending; < US$1 is a source target only.')],[1.45,3.25,1.8])
    h('8. Methodology and Implementation Plan')
    h('8.1 Overall Methodology',2)
    p('Proceed through requirements, citation and dataset verification, architecture design, baseline implementation, candidate-role discovery, cloud integration, security review, and repeatable evaluation. Keep Phase-I design artifacts distinct from Phase-II executable evidence. Compare alternative pipeline stages through controlled ablations, recording code revision, parameters, environment, dataset version, and random seed.')
    h('8.2 Tools and Technologies',2)
    caption('Table 7. Proposed toolchain.')
    table(['Tool / technology','Purpose','Version / configuration'],[
        ('Python; sparse data libraries','Parsing, clustering, role mining, and evaluation.','Exact dependencies to be pinned for Phase II.'),
        ('K-Means / hierarchical clustering','Attribute-assisted candidates where attributes are real.','Distance, encoding, cluster count, and seed to be documented.'),
        ('HTML / JavaScript or React','Administrative dashboard proposed in the source.','Frontend framework version not specified.'),
        ('AWS services and SDK','APIs, storage, authentication, and scoped checks.','Region and SDK/runtime versions not supplied.'),
        ('Git and GitHub','Version control, reviews, and release evidence.','Source repository and proposed team branches.'),
        ('Local Python / Google Colab','Offline preparation and repeatable experiments.','Record hardware, runtime, and package lock file.')],[1.7,2.6,2.2])
    h('8.3 Step-by-Step Implementation Plan',2)
    for s in [
        'Requirements and success criteria: define legitimate access, reference policies, supported mappings, and metric thresholds.',
        'Dataset/input preparation: acquire licensed inputs, verify counts, label synthetic additions, and save a provenance manifest.',
        'Architecture and AWS setup: select region, define private storage/authentication, and prepare repeatable infrastructure configuration.',
        'Application/backend/model development: implement parsing, a baseline, candidate roles, reconstruction, WSC, and the dashboard.',
        'Integration and security: map eligible permissions to AWS actions/resources, run scoped checks, and add approval/version controls.',
        'Testing and performance evaluation: run functional, security, reconstruction, ablation, and cold/warm workload experiments.',
        'Deployment, monitoring, documentation: deploy in a test account, capture dated evidence, monitor cost, and publish a reviewed release.'
    ]: p(s,'List Number')
    h('8.4 Evaluation Metrics',2)
    p('WSC = w_r|R| + w_u|UA| + w_p|PA| + w_h|RH| + w_d|DUPA|, following the source’s Molloy et al. (2009) formulation. R denotes roles, UA user–role assignments, PA role–permission assignments, RH hierarchy edges, and DUPA direct user–permission assignments. Declare every weight and count the baseline under the same representation. WSC reduction (%) = 100 × (WSC_baseline − WSC_candidate) / WSC_baseline, for a positive baseline value.')
    p('For input grants A and reconstructed grants B, report missing = |A − B| and extra = |B − A|. Exact reconstruction requires both to be zero. Also report role count, assignment count, candidate-search size, runtime and peak memory. Silhouette score applies only to the documented clustering representation and distance. Measure end-to-end latency over a specified dataset size and repetitions, reporting cold/warm conditions and percentiles. Policy-check findings, regression failures, and administrator task time are separate measures; no single metric proves universal least privilege.')
    h('9. Implementation / Development')
    p('This section reflects the implementation status recorded in the supplied Phase-I PDF. It does not infer completed cloud deployment from architecture drawings or planned features.')
    h('9.1 Implemented Modules',2)
    caption('Table 8. Documented module and deliverable status.')
    table(['Module / deliverable','Implementation / evidence','Status'],[
        ('Problem, objectives, novelty','Source PDF pp. 19–21; consolidated in Sections 1, 3, 4.','Documented in Phase I.'),
        ('Literature and gaps','Source pp. 1–18; duplicate and unresolved citations identified in Section 2.','Documented; literature requirement incomplete.'),
        ('AWS and system design','Source pp. 21–23; Figures 1–2 redrawn in this report.','Design completed; deployment unverified.'),
        ('Dataset/preprocessing design','Source pp. 23–24; no input manifest supplied.','Design documented; data unverified.'),
        ('Attribute clustering and sparse filtering','Source p. 26 states IN PROGRESS. No test artifact supplied.','In progress, as reported.'),
        ('IAM transpiler / Access Analyzer integration','Source p. 26 states PLANNED.','Planned.'),
        ('Repository/team workflow','Source p. 26 names repository and branches. No commit/PR evidence supplied.','Reported active; evidence pending.')],[1.65,3.2,1.65])
    h('9.2 Screenshots and Configuration Evidence',2)
    p('The PDF contains two architecture illustrations on page 22, redrawn here as Figures 1 and 2. It contains no dated application screenshots, console deployment screenshots, API responses, policy-check output, billing reports, or test logs. Those evidence items remain pending. Future evidence should show a timestamp, dataset identifier, environment, and relevant code revision while concealing credentials and personal information.')
    h('9.3 GitHub Repository',2)
    p('Repository supplied in the PDF: https://github.com/Ansh-bit948/AWS-project. The template suggests the naming convention ProjectName_Cloud_Project_2026; the supplied repository name differs. No repository rename or remote change is implied by this report. Source layout, commit hashes, releases, and pull-request links were not included in the PDF.')
    p('A proposed organisation is README.md and dependency/configuration files at the root; docs/ for reports and literature; architecture/ for diagrams; dataset/ for licensed inputs and manifests; src/ for application and cloud code; results/ for measurement artifacts; and presentation/ for review slides. This is a documentation plan rather than a verified inventory of the remote repository.')
    h('10. Testing, Results and Observations')
    h('10.1 Test Plan and Test Cases',2)
    caption('Table 9. Planned functional, security, and performance tests.')
    tests=[
        ('T01','Input validation','Valid and malformed CSV/JSON inputs.','Accept valid data; explain invalid fields.'),
        ('T02','Grant reconstruction','Known UPA matrix and mined roles.','Zero missing and zero extra grants in exact mode.'),
        ('T03','WSC comparison','Identical input and declared weights.','Reproducible baseline/candidate counts and reduction.'),
        ('T04','Authentication','Missing, expired, and valid tokens.','Reject invalid access; accept authorised calls.'),
        ('T05','Policy mapping/checks','Supported and unsupported action mappings.','Reject unresolved mappings and record validation findings.'),
        ('T06','Approval/rollback','Approved change and induced regression.','Versioned change; recover previous approved state.'),
        ('T07','Latency and scale','Fixed dataset sizes; cold/warm repeated runs.','Report duration distribution, memory, failures, and environment.'),
        ('T08','Cost and operations','Controlled demo and service usage records.','Document actual usage; verify alert delivery separately.')]
    table(['Test / scenario','Input / expected result','Observed / verdict'],[(i+' '+s,inp+' '+exp,'Not run in supplied evidence; pending.') for i,s,inp,exp in tests],[1.35,3.55,1.6])
    h('10.2 Results',2)
    p('The supplied PDF contains numerical statements in its abstract but no result tables, experiment logs, dataset manifest, or reproducible runs supporting them. Its later status table identifies core implementation as in progress or planned. These statements are therefore recorded below as unverified source claims, not validated outcomes.')
    caption('Table 10. Source numerical claims and evidence status.')
    table(['Source claim / target','Available evidence','Reporting status'],[
        ('Over 42% WSC reduction','No baseline weights, raw counts, code run, or measurement artifact.','Unverified claim; ≥ 40% remains a target.'),
        ('68% administration overhead reduction','No administrator protocol, sample, timing data, or baseline.','Unverified claim; study required.'),
        ('100% wildcard-path elimination','No generated policy corpus or checks.','Unverified claim; limited test objective only.'),
        ('112,120 tuples / 4,930 users / 560 permissions','Dataset specification only.','Counts require an input manifest.'),
        ('15 benchmark suites / 15 papers','No benchmark suite inventory; five numbered literature entries repeat.','Neither count is established as distinct evaluated studies.'),
        ('< 2.5 seconds and < US$1','Objectives without workload or measured bill.','Targets only; not evaluated.')],[1.8,3.05,1.65])
    h('10.3 Baseline Comparison',2)
    p('Not yet evaluated in the supplied evidence. The planned baseline is the original direct-grant representation or a documented simple role construction, measured with the same WSC weights and identical source grants. Compare baseline, structural-only mining, attribute-assisted mining, and optional graph processing under equal runtime constraints. Report all runs, including cases with no WSC improvement, and verify reconstruction before interpreting compression.')
    h('10.4 Results Summary',2)
    p('Phase I establishes a problem statement, architecture, objectives, and implementation plan. None of the six runtime objectives is shown as experimentally achieved by the supplied PDF. Attribute mining is reported in progress; policy synthesis, scoped validation, authenticated deployment, latency, and cost measurements remain pending. Completion of design documents should be reported separately from successful functional execution.')
    h('11. Discussion, Limitations and Risk Analysis')
    h('11.1 Discussion',2)
    p('The integrated design is a plausible way to connect permission analytics with administrator decisions. Its evaluation must establish both structural fidelity and operational usefulness. A smaller role set may simplify administration but can still encode unnecessary existing privileges. Domain labels can assist interpretation only when they come from genuine, authorised records. Cloud validation checks complement, rather than replace, application authorisation tests and human judgement.')
    h('11.2 Limitations',2)
    bullets(['The source literature has repeated entries and incomplete publication records; 15 distinct verified research papers are not established.',
             'The dataset specification is not accompanied by files, license text, manifests, or a synthetic-data generator.',
             'Static benchmark assignments do not provide timestamps, real job attributes, or approval outcomes.',
             'Clustering, graph pruning, policy generation, and deployment have no reproducible runtime evidence in the PDF.',
             'Sparse logs can omit legitimate rare activity; structural anomalies do not establish threats.',
             'Mapping a university ERP permission to an AWS IAM permission requires an explicit service/resource model.',
             'The combined Neptune architecture and unverified Free Tier assumptions do not establish zero idle cost.',
             'Latency, scalability, administrator effort, and security generalisation require controlled experiments.'])
    h('11.3 Risks and Mitigation',2)
    caption('Table 11. Project risks and planned mitigation.')
    table(['Risk','Likelihood / impact','Mitigation / residual risk'],[
        ('Unexpected AWS cost','Medium / high','Estimate per region; set alerts; bound concurrency and retention; cost optional Neptune separately. Alerts can arrive after spending.'),
        ('Credentials or misconfiguration','Medium / high','Scoped roles, no keys in Git, private storage, sanitised logs, and reviewed policies. Residual risk requires account testing.'),
        ('Dataset provenance or privacy','High / high','Verify license, source hashes, authorisation, and feature origins; anonymise records. Real-domain access remains unavailable.'),
        ('Incorrect privilege removal','Medium / high','Review anomalies; preserve versions; regression-test legitimate workflows; rollback. Incomplete logs remain a risk.'),
        ('Unscalable analysis','Medium / medium','Sparse processing, bounded jobs, measured input limits, asynchronous design where needed. Large matrices remain untested.'),
        ('Unsupported research/results','High / high','Resolve citations, replace duplicate papers, retain raw experiment evidence, and label targets. Research completeness remains pending.')],[1.5,1.1,3.9])
    h('12. Team Contribution and GitHub Management')
    p('Names, registrations, paper allocation, and service ownership follow the source PDF. The table distinguishes recorded design allocation from work still planned; it does not certify unprovided commits or independently authored reviews.')
    caption('Table 12. Team allocation and contribution status.')
    table(['Activity','Ansh Goel\n24BIT0258','Agastya Pradhan\n24BIT0187','Varun Verma\n24BIT0208'],[
        ('Literature','P1–P5 source section.','P6–P10 section repeats P1–P5.','P11–P15 source section; citations unresolved.'),
        ('Gap / novelty','Attribute-guided mining direction.','Structural / governance direction.','Policy integration direction.'),
        ('Frontend','Individual allocation not stated.','Individual allocation not stated.','Individual allocation not stated.'),
        ('Backend / database','Input/artifact service planning.','DynamoDB design owner.','Lambda/API planning owner.'),
        ('AWS / security','S3, graph extension, CloudWatch.','Cognito, SNS.','IAM/Access Analyzer and API.'),
        ('Dataset / analytics','Clustering branch proposed.','Graph branch proposed.','Policy mapping branch proposed.'),
        ('Testing / results','Planned; no individual result evidence.','Planned; no individual result evidence.','Planned; no individual result evidence.'),
        ('Documentation','Named in source Phase-I report.','Named in source Phase-I report.','Named in source Phase-I report.')],[1.15,1.78,1.78,1.79])
    h('12.1 Branch and Pull Request Workflow',2)
    p('The source proposes main for releases, develop for integration, and feature/ansh-clustering, feature/agastya-graph, and feature/varun-iam for individual changes. Each change should identify its purpose, tests, and evidence in a pull request; another member reviews it before merge. Resolve conflicts on the feature branch, rerun relevant checks, and tag the final reviewed release. These practices are a stated workflow, not verified repository history.')
    h('12.2 Contribution Evidence',2)
    p('The source identifies the repository owner through Ansh-bit948 but supplies no individual profile URLs for the other members, commit hashes, pull-request links, review records, or release tags. The template’s expectation of roughly 20–30 meaningful commits per student, at least two pull requests, review participation, and weekly activity cannot be confirmed. Attach an evidence register per student before final submission; commit counts alone do not prove substantive contribution.')
    h('13. Conclusion and Future Enhancements')
    h('13.1 Conclusion',2)
    p('This Phase-I report consolidates the proposed permission-optimisation framework into the supplied cloud-project structure. The design combines sparse access analysis, role discovery, WSC comparison, an authenticated AWS workflow, and controlled policy review. The established outputs are design and planning artifacts, with clustering reported in progress and policy integration planned. No measured compression, runtime, operational cost, or broad security guarantee is established by the PDF. The next phase must turn the design into reproducible code and evidence while resolving literature duplication, dataset provenance, and explicit application-to-cloud mappings.')
    h('13.2 Future Enhancements',2)
    bullets(['Complete a survey of 15 distinct, verifiable studies and independent student analyses.',
             'Publish an authorised dataset inventory with true counts, licenses, hashes, and clearly labelled synthetic fields.',
             'Implement a bounded baseline and role-mining pipeline with reconstruction and WSC reporting.',
             'Add approval-backed policy mapping, scoped checks, regression tests, and rollback.',
             'Deploy a small controlled test workload and measure latency, reliability, and account-specific cost.',
             'Evaluate interpretability and review effort through an explicit administrator study.',
             'Investigate incremental updates and optional graph processing only after demonstrating the core workflow.'])
    references()
    appendices()
    doc.core_properties.title=TITLE
    doc.core_properties.subject='BCSE355L Cloud Architecture Design — Phase-I report'
    doc.core_properties.author='; '.join(n for n,_ in TEAM)
    doc.core_properties.keywords='AWS, role mining, university ERP, WSC, Phase I'
    doc.core_properties.comments='Prepared from the user-supplied report PDF and Word template; unverified source claims are identified.'
    doc.save(TARGET)
    assert 200 <= len(ABSTRACT.split()) <= 300, len(ABSTRACT.split())
    print('Saved',TARGET)
    print('Abstract words:',len(ABSTRACT.split()))
    print('Paragraphs:',len(doc.paragraphs),'Tables:',len(doc.tables),'Figures:',len(doc.inline_shapes))

def references():
    h('References')
    p('Research citations below are transcribed from the supplied PDF in APA-style order. The first five entries are source-provided bibliographic records; their metadata has not been independently audited. Entries with incomplete or ambiguous source attribution are listed separately so they are not presented as verified publications. Duplicate P6–P10 entries are not repeated as additional publications.')
    for s in [
        'Anderer, S., Kreppein, D., Scheuermann, B., & Mostaghim, S. (2020). The addRole-EA: A new evolutionary algorithm for the role mining problem. In Proceedings of the 12th International Joint Conference on Computational Intelligence (pp. 155–166). SCITEPRESS. (P3 / P8; details as supplied.)',
        'Hamid, N. A., Ahmad, R., & Selamat, S. R. (2017). Recent trends in role mining algorithms for role-based access control: A systematic review. World Applied Sciences Journal, 35(7), 1054–1058. (P4 / P9; details as supplied.)',
        'Lu, H., Vaidya, J., & Atluri, V. (2014). An optimization framework for role mining. Journal of Computer Security, 22(1), 1–31. (P2 / P7; details as supplied.)',
        'Mollakuqe, E., & Dimitrova, V. (2024). Comparative analysis of identity management, access control, and authorization practices in public and private universities. Open Research Europe, 4, Article 23. (P5 / P10; details as supplied.)',
        'Molloy, I., Li, N., Li, T., Mao, Z., Wang, Q., & Lobo, J. (2009). Evaluating role mining algorithms. In Proceedings of the 14th ACM Symposium on Access Control Models and Technologies (pp. 95–104). ACM. (P1 / P6; details as supplied.)'
    ]: p(s)
    h('Unresolved source references',2)
    for a in PAPERS[5:]: p(f'{a[0]} — {a[2]}. Attribution in source: {a[1]}. Exact author list, title, venue, year, and DOI/stable publication URL require verification. This entry is retained for source traceability and is not counted as a verified reference.')
    h('Official AWS documentation',2)
    refs=[
        ('a','CloudTrail concepts','https://docs.aws.amazon.com/awscloudtrail/latest/userguide/cloudtrail-concepts.html'),
        ('b','Using AWS Identity and Access Management Access Analyzer','https://docs.aws.amazon.com/IAM/latest/UserGuide/what-is-access-analyzer.html'),
        ('c','Amazon Neptune pricing','https://aws.amazon.com/neptune/pricing/'),
        ('d','AWS Budgets pricing','https://aws.amazon.com/aws-cost-management/aws-budgets/pricing/')]
    for ident,title,url in refs: p(f'Amazon Web Services. (n.d.-{ident}). {title}. Retrieved October 15, 2026, from {url}')
    h('Project source documents',2)
    p('Goel, A., Pradhan, A., & Verma, V. (2026, July 31). Intelligent permission optimization framework for university ERP systems using role mining analytics [Unpublished Phase-I project report; user-supplied aws-project.pdf].')
    p('AWS Cloud Project Report Template. (n.d.). [User-supplied Word template; AWS_Cloud_Project_Report_Template.docx].')

def appendices():
    h('Appendix A – Project Timeline')
    p('The source records a Phase-I submission date of 31 July 2026 but no task-level start/end dates. Phase-II timing below is relative and proposed, avoiding invented historical milestones.')
    caption('Table 13. Source milestones and proposed Phase-II sequence.')
    table(['Phase / task','Timing','Owner(s)','Status / evidence'],[
        ('Literature, gap, objectives','Phase I; recorded by 31/07/2026','All members','Source documents present; unique-paper review incomplete.'),
        ('Architecture and dataset planning','Phase I; recorded by 31/07/2026','All members','Source pp. 21–25; design only.'),
        ('Citation and dataset audit','Phase-II week 1 (proposed)','All members','Pending.'),
        ('Baseline and clustering','Phase-II weeks 2–3 (proposed)','Ansh; team review','Source reports clustering in progress.'),
        ('Graph / storage evaluation','Phase-II weeks 3–4 (proposed)','Agastya; Ansh','Pending; Neptune optional.'),
        ('Policy mapping / checks / API','Phase-II weeks 4–5 (proposed)','Varun; team review','Planned in source.'),
        ('Integration and measured tests','Phase-II weeks 6–7 (proposed)','All members','Pending.'),
        ('Final evidence and release','Phase-II week 8 (proposed)','All members','Pending.')],[1.65,1.65,1.15,2.05])
    h('Appendix B – Review Checklist')
    caption('Table 14. Report preparation and evidence checklist.')
    table(['Requirement','Status'],[
        ('Abstract: 200–300 words with status','Included; quantitative claims separated from verified outcomes.'),
        ('15 distinct research papers and comparison','Incomplete in source; duplicate and unresolved entries disclosed.'),
        ('Independent student gap analysis','Source allocations recorded; duplicate analysis requires replacement.'),
        ('4–6 measurable objectives','Six objectives with planned metrics included.'),
        ('Concise, defensible novelty summary','Included as proposed integration contribution.'),
        ('Two architecture diagrams','Included and explained; redrawn from source design.'),
        ('Dataset description and provenance','Specification included; source files, counts, license, and URL pending.'),
        ('AWS services and cost','Planned status documented; numeric estimate pending.'),
        ('Implementation/tests/results','Status and test plan included; runtime evidence pending.'),
        ('GitHub and individual contributions','Repository/workflow retained; verifiable commit/PR links pending.'),
        ('APA references and citations','Source citations normalised; unresolved records identified.'),
        ('Credentials and confidential data','Report excludes credentials; future evidence needs sanitisation.'),
        ('Font, headings, spacing, margins, pages','Times New Roman 12 pt; 14–16 pt headings; 1.5 spacing; 1-inch margins; centred page numbers.')],[2.4,4.1])
    h('Appendix C – Source Reconciliation and Evidence Register')
    p('This appendix records material corrections made while transferring the PDF into the template. It preserves the original Phase-I date and student identities, uses the body’s consistent 1–13 chapter sequence, removes template instructional prompts, and separates reported progress from measured evidence. It does not certify that the current software repository matches the historical PDF.')
    caption('Table 15. Source content reconciliation.')
    table(['Source item','Treatment in this report'],[
        ('P6–P10 repeat P1–P5','Retained as source identifiers; disclosed as duplicates.'),
        ('P11–P15 incomplete citations','Retained as source summaries with unresolved reference status.'),
        ('42%, 68%, 100% achieved claims','Listed as unverified because later implementation status and missing evidence do not substantiate them.'),
        ('Dataset scale and feature claims','Recorded as specifications awaiting a manifest and preprocessing explanation.'),
        ('Universal Zelkova proof','Reframed as scoped Access Analyzer checks plus regression tests and review.'),
        ('CloudTrail captures all ERP actions','Clarified AWS activity scope; application audit export required.'),
        ('Zero idle cost / fixed free allowances','Replaced with account-specific estimation requirements and optional Neptune costing.'),
        ('US$1 hard cap','Described as an alert threshold, not a spending cap.'),
        ('Original diagrams, PDF p. 22','Extracted into companion assets and redrawn at readable resolution.'),
        ('Programme/year and team number','Marked not specified; no identities or administrative values invented.')],[2,4.5])
    p('Evidence needed for the final implementation report: exact publication records; independently authored paper assignments; dataset files and manifest; licenses; preprocessing scripts; code revision; unit/integration and regression outputs; architecture/deployment configuration; dated screenshots; validation findings; cost worksheet and usage report; per-student commits and reviewed pull requests.')

if __name__=='__main__': main()
